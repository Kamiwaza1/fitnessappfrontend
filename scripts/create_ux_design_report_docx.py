from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "ux-design-report.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
MUTED = RGBColor(96, 112, 128)
LIGHT_GRAY = "F2F4F7"
LIGHT_BLUE = "E8EEF5"
CALLOUT = "F4F6F9"
GREEN = RGBColor(22, 163, 74)
WHITE = RGBColor(255, 255, 255)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tbl_cell_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.allow_autofit = False
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = Inches(width)
            tc_pr = row.cells[index]._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(round(width * 1440)))

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(round(width * 1440) for width in widths)))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")

    tbl_grid = table._tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(round(width * 1440)))
        tbl_grid.append(grid_col)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_para(doc, text="", style=None, before=0, after=6, line=1.10, align=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(text)
        set_run_font(run, color=INK)
    return p


def add_heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    fmt = p.paragraph_format
    if level == 1:
        fmt.space_before = Pt(16)
        fmt.space_after = Pt(8)
        size = 16
        color = BLUE
    elif level == 2:
        fmt.space_before = Pt(12)
        fmt.space_after = Pt(6)
        size = 13
        color = BLUE
    else:
        fmt.space_before = Pt(8)
        fmt.space_after = Pt(4)
        size = 12
        color = DARK_BLUE
    fmt.line_spacing = 1.10
    for run in p.runs:
        set_run_font(run, size=size, color=color, bold=True)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_run_font(run, color=INK)


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_run_font(run, color=INK)


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_width(table, [6.5])
    set_cell_margins(table, top=120, bottom=120, start=160, end=160)
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=11, color=DARK_BLUE, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.10
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_key_value_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    set_table_width(table, [1.875, 4.625])
    set_cell_margins(table)
    for i, (label, value) in enumerate(rows):
        label_cell, value_cell = table.rows[i].cells
        set_cell_shading(label_cell, LIGHT_GRAY)
        label_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        value_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        label_run = label_cell.paragraphs[0].add_run(label)
        value_run = value_cell.paragraphs[0].add_run(value)
        set_run_font(label_run, size=10, color=DARK_BLUE, bold=True)
        set_run_font(value_run, size=10, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_matrix(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table, widths)
    set_cell_margins(table)
    repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_BLUE)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        run = cell.paragraphs[0].add_run(header)
        set_run_font(run, size=9.5, color=DARK_BLUE, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            run = cells[index].paragraphs[0].add_run(value)
            set_run_font(run, size=9.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10

    header = section.header.paragraphs[0]
    header.text = "Dark Lift UX Design Report"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.text = "Prepared for FitnessAppFrontend"
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = MUTED


def add_title_page(doc):
    add_para(doc, "UX DESIGN REPORT", before=40, after=8)
    p = doc.paragraphs[-1]
    p.runs[0].font.size = Pt(23)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = RGBColor(0, 0, 0)

    p = add_para(doc, "Dark Lift Fitness App", after=18)
    p.runs[0].font.size = Pt(15)
    p.runs[0].font.color.rgb = MUTED

    add_key_value_table(
        doc,
        [
            ("Product", "Dark Lift fitness tracking web app"),
            ("Repository", "FitnessAppFrontend / fitnessappfrontend"),
            ("Prepared", date.today().strftime("%B %-d, %Y")),
            ("Scope", "Landing, authentication, dashboard, workouts, nutrition, progress, settings, and navigation"),
            ("Preset", "Standard business brief"),
        ],
    )

    add_callout(
        doc,
        "Executive takeaway",
        "Dark Lift has a strong feature foundation and a memorable athletic brand. The highest-impact UX work is to make core flows calmer, clearer, and faster on mobile while improving empty states, accessibility, and feedback patterns.",
    )


def build_document():
    doc = Document()
    configure_document(doc)
    add_title_page(doc)

    add_heading(doc, "Product Scope Reviewed", 1)
    for item in (
        "Landing, login, registration, password reset",
        "Authenticated app shell and navigation",
        "Dashboard",
        "Workout planning and active workout session",
        "Live workout entry points",
        "Exercise library",
        "Workout history",
        "Nutrition diary",
        "Progress and insights",
        "Settings",
    ):
        add_bullet(doc, item)

    add_heading(doc, "Target Users", 1)
    for item in (
        "Regular gym users tracking planned workouts and performance",
        "Fitness beginners who need guided structure and simple progress feedback",
        "Users tracking nutrition alongside training",
        "Small groups or partners using live workout rooms",
    ):
        add_bullet(doc, item)

    add_heading(doc, "Core User Journey Findings", 1)

    journeys = [
        (
            "New User Onboarding",
            "The app introduces the Dark Lift brand through a bold landing and authentication experience. The login form is clear, with direct links for account creation and password recovery.",
            [
                "Brand recognition is immediate.",
                "Sign-in flow is simple and familiar.",
                "Form validation and error messages are present.",
            ],
            [
                "The app does not clearly preview the product value before login.",
                "The authentication experience uses a more decorative visual style than the operational app.",
            ],
            [
                "Add a concise product promise to the landing hero, focused on tracking workouts, nutrition, and progress.",
                "Align typography, spacing, and button treatments with the logged-in app.",
            ],
        ),
        (
            "Dashboard Review",
            "The dashboard provides an overview of daily calories, protein, workout status, trends, and recommendations.",
            [
                "It answers common questions quickly: what is planned today, how calories are tracking, and what to do next.",
                "Circular calorie and protein indicators provide fast visual feedback.",
                "Recommendations create a clear path back into workouts.",
            ],
            [
                "Several dashboard elements compete visually for attention.",
                "Date rendering relies on client-side time and hydration suppression.",
                "Empty or loading states could provide more guidance.",
            ],
            [
                "Prioritize one primary dashboard action, such as starting today's workout.",
                "Use consistent hierarchy: key metrics first, charts second, recommendations third.",
                "Provide actionable empty states such as creating a workout or logging the first meal.",
            ],
        ),
        (
            "Workout Creation and Session Tracking",
            "The workout area supports creating workouts, editing exercises, starting active sessions, joining live rooms, and finishing sessions.",
            [
                "The workflow covers the full training loop from planning to completion.",
                "The start confirmation modal helps prevent accidental session replacement.",
                "Active workout state appears persistent and globally visible.",
            ],
            [
                "The workout composer is dense on mobile.",
                "Destructive actions use browser alerts in some places.",
                "Start, edit, delete, and live workout actions need clearer grouping.",
            ],
            [
                "Replace browser alerts with app-native dialogs or toast messages.",
                "Split workout creation into workout details, exercise list, and review sections.",
                "Keep primary, secondary, and destructive action hierarchy consistent.",
            ],
        ),
        (
            "Nutrition Logging",
            "Nutrition includes manual entry, OCR, and food database search, which gives users multiple ways to reduce logging friction.",
            [
                "Multiple input modes support different user preferences.",
                "Macro fields are explicit and structured.",
                "Meal category selection supports diary organization.",
            ],
            [
                "Manual, OCR, and database modes may feel equally weighted even though manual should be the reliable fallback.",
                "OCR results require strong confirmation before saving.",
                "Database results need clearer serving size and brand information.",
            ],
            [
                "Rename OCR to Scan label for non-technical users.",
                "Show a review state after scan extraction.",
                "Add recent or frequent foods to speed up repeat logging.",
            ],
        ),
        (
            "Progress and Insights",
            "Progress shows workout categories, duration trends, and personal records.",
            [
                "Charts give the product a useful analytical layer.",
                "Time-range controls are familiar for progress pages.",
                "Personal records are motivational and worth expanding.",
            ],
            [
                "Some controls look available before their behavior is clear.",
                "Empty personal records currently feel like a dead end.",
                "The chart palette is heavily green, reducing category distinction.",
            ],
            [
                "Disable or wire up time-range controls consistently.",
                "Improve empty records with next-step guidance.",
                "Add summary metrics above charts: total workouts, volume, average duration, and streak.",
            ],
        ),
    ]

    for title, summary, strengths, issues, recommendations in journeys:
        add_heading(doc, title, 2)
        add_para(doc, summary)
        add_heading(doc, "Strengths", 3)
        for item in strengths:
            add_bullet(doc, item)
        add_heading(doc, "Issues", 3)
        for item in issues:
            add_bullet(doc, item)
        add_heading(doc, "Recommendations", 3)
        for item in recommendations:
            add_bullet(doc, item)

    add_heading(doc, "Navigation and Information Architecture", 1)
    add_para(
        doc,
        "The sidebar structure is logical and covers the major product areas: Dashboard, Workouts, Exercise Library, Live Workout, History, Nutrition, Progress, and Settings.",
    )
    for item in (
        "Desktop navigation is persistent and predictable.",
        "Mobile bottom navigation keeps key areas accessible.",
        "Active live workout state is surfaced globally, which is useful.",
        "Mobile navigation has many items, making labels small and horizontally scrollable.",
        "Consider a five-item mobile nav: Dashboard, Workouts, Nutrition, Progress, More.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "Visual Design Assessment", 1)
    add_para(
        doc,
        "Dark Lift has a clear brand: bold, athletic, high-contrast, and green-accented. The app uses strong headings, icon-led navigation, charts, badges, and status panels to create energy.",
    )
    for item in (
        "Keep green for primary actions and key highlights.",
        "Add neutral, amber, blue, and red semantic colors for secondary chart series and state messaging.",
        "Standardize border radius around 8px for most controls and cards.",
        "Reduce decorative styling inside data-heavy pages.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "Accessibility Review", 1)
    for item in (
        "Add aria-hidden=\"true\" to decorative Material icon spans.",
        "Ensure all icon-only or compact controls have accessible names.",
        "Verify contrast for primary green on light and dark backgrounds.",
        "Use app-native dialogs and toasts with clear focus management.",
        "Add keyboard testing for composer, modal, bottom nav, and live workout flows.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "Responsive UX Review", 1)
    for item in (
        "Use stacked form sections with sticky save actions on mobile.",
        "Collapse charts into summary cards first, with expandable detail.",
        "Replace wide tables with mobile-friendly list rows.",
        "Test all primary flows at 375px width.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "Content and Microcopy", 1)
    for item in (
        "Replace OCR with Scan label in the UI.",
        "Use empty states that pair status with action.",
        "Keep motivational copy short and tied to the user's task.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "Priority Recommendations", 1)
    add_matrix(
        doc,
        ["Priority", "Recommendation", "UX Impact"],
        [
            ("High", "Fix hydration mismatch in root layout font loading.", "Improves reliability and removes distracting runtime errors."),
            ("High", "Simplify mobile navigation.", "Reduces crowding and improves discoverability on small screens."),
            ("High", "Improve empty states.", "Guides users toward productive first actions."),
            ("High", "Replace browser alerts with in-app feedback.", "Creates consistent, accessible interaction patterns."),
            ("High", "Review icons, focus states, and contrast.", "Improves accessibility and keyboard usability."),
            ("Medium", "Refine dashboard hierarchy around one daily action.", "Makes the dashboard faster to scan and act on."),
            ("Medium", "Improve workout composer structure for mobile.", "Reduces form fatigue during workout creation."),
            ("Medium", "Add review step for scanned nutrition results.", "Prevents incorrect nutrition data from being saved."),
            ("Medium", "Add recent or frequent foods.", "Speeds up repeat nutrition logging."),
            ("Medium", "Expand chart palette beyond green variations.", "Improves data distinction and readability."),
        ],
        [1.0, 3.0, 2.5],
    )

    add_heading(doc, "Suggested UX Success Metrics", 1)
    for item in (
        "Workout creation completion rate",
        "Active workout start-to-finish completion rate",
        "Average time to log a meal",
        "Repeat nutrition logging rate",
        "Dashboard recommendation click-through rate",
        "Mobile navigation mis-tap or backtrack rate",
        "Form validation error rate",
        "Weekly active users completing at least one workout",
    ):
        add_bullet(doc, item)

    add_heading(doc, "Conclusion", 1)
    add_para(
        doc,
        "Dark Lift already has the foundation of a useful fitness operating system: planning, active sessions, live workouts, nutrition, and progress are all represented. The next UX step is not adding more features, but making the existing flows feel calmer, clearer, and faster. Prioritizing mobile navigation, dashboard hierarchy, accessible feedback patterns, and better empty states will make the app feel significantly more polished with a controlled implementation effort.",
    )

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build_document()
    print(path)
