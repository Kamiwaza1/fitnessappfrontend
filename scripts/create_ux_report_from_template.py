from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path("/Users/zekilok/Downloads/UX Design Report.docx")
OUT = ROOT / "docs" / "dark-lift-ux-design-report-template.docx"
SCREENSHOT_DIR = ROOT / "cypress" / "screenshots" / "ux-report-screenshots.cy.js" / "ux-report"


def clear_document_body(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_run(run, *, bold=False, italic=False, size=14) -> None:
    run.bold = bold
    run.italic = italic
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)


def add_rich_line(doc: Document, parts, *, size=13, after=5) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    for text, bold, italic in parts:
        run = p.add_run(text)
        set_run(run, bold=bold, italic=italic, size=size)


def add_body(doc: Document, text: str, *, after=8) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text)
    set_run(run, size=13)


def add_heading(doc: Document, text: str, *, size=14, after=8) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    set_run(run, bold=True, size=size)


def add_screenshot(doc: Document, filename: str, width=5.35) -> None:
    path = SCREENSHOT_DIR / filename
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(7)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))


def add_page_break(doc: Document) -> None:
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_mistake(
    doc: Document,
    number: int,
    title: str,
    location: str,
    description: str,
    screenshot: str,
    heuristic_number: str,
    heuristic_name: str,
    heuristic_quote: str,
    heuristic_explanation: str,
    proposed_fix: str,
    *,
    page_break=False,
) -> None:
    if page_break:
        add_page_break(doc)
    add_heading(doc, f"Mistake {number} — {title}", size=14, after=9)
    add_rich_line(doc, [("Location: ", True, False), (location, False, False)], size=13, after=7)
    add_body(doc, description, after=8)
    add_screenshot(doc, screenshot)
    add_rich_line(
        doc,
        [
            ("Heuristic violated: ", True, False),
            (f"{heuristic_number} — {heuristic_name}\n", False, False),
            (f'"{heuristic_quote}" ', False, True),
            (heuristic_explanation, False, False),
        ],
        size=13,
        after=7,
    )
    add_rich_line(doc, [("Proposed fix: ", True, False), (proposed_fix, False, False)], size=13, after=8)


def build() -> Path:
    doc = Document(str(TEMPLATE))
    clear_document_body(doc)

    add_heading(doc, "UX Design Report: Dark Lift Heuristic Evaluation", size=14, after=12)
    add_rich_line(
        doc,
        [
            ("Tester: ", True, False),
            ("Zekilok / Codex\n", False, False),
            ("Date: ", True, False),
            (date.today().strftime("%B %-d, %Y") + "\n", False, False),
            ("Application: ", True, False),
            ("Dark Lift — Fitness Tracking Web App\n", False, False),
            ("Method: ", True, False),
            ("Heuristic evaluation against Jakob Nielsen's 10 Usability Heuristics, using live frontend screenshots from localhost:3000.", False, False),
        ],
        size=13,
        after=28,
    )

    add_mistake(
        doc,
        1,
        "Dashboard Shows Long-Lived Loading and Syncing States",
        "Dashboard page — daily calorie/protein cards, today's goal panel, recommendations area",
        "The dashboard screenshot shows key data areas stuck in transitional states: daily calories and protein display 0, the goal card says \"SYNCING...\", and recommendations show \"Loading recommendations...\". Because these elements are primary dashboard content, users may not know whether the app has no data, the network is still loading, or something failed silently.",
        "dashboard-overview.png",
        "#1",
        "Visibility of System Status",
        "The system should always keep users informed about what is going on, through appropriate feedback within a reasonable amount of time.",
        "Dashboard loading states need to resolve into loaded, empty, or error states so users can understand the system status.",
        "Add explicit empty and error states for each dashboard section. For example: \"No meals logged today\" with a Log Food action, \"No workout scheduled\" with a Plan Workout action, and \"Recommendations unavailable\" with a retry affordance. Avoid showing indefinite loading copy after the request has failed or returned empty data.",
        page_break=False,
    )

    add_mistake(
        doc,
        2,
        "Workout Composer Is Dense and Partially Clipped on Desktop",
        "Workouts page — Add Workout composer",
        "The workout composer opens as a large two-column form with session details on the left and exercises on the right. In the screenshot, key controls such as Add Workout, Close, Add Exercise, and Save Workout are pushed to the far right or bottom edge. This makes the creation flow feel cramped and increases the risk that users miss the available actions.",
        "workout-composer.png",
        "#8",
        "Aesthetic and Minimalist Design",
        "Interfaces should not contain information which is irrelevant or rarely needed.",
        "Dense forms make it harder to scan the next required action and increase cognitive load during workout creation.",
        "Break the composer into clearer sections: Workout Details, Exercises, and Review. Keep the primary Save Workout action visible in a consistent footer, and make secondary actions such as Add Exercise and Close easier to reach without relying on edge placement.",
    )

    add_mistake(
        doc,
        3,
        "Mobile Navigation Is Crowded and Labels Are Truncated",
        "Mobile dashboard — bottom navigation and active viewport",
        "On the mobile screenshot, the bottom navigation has many destinations competing for limited width. Labels such as Exercise Library and Live Workout are truncated, and a visible issue badge overlaps the navigation area. This makes navigation harder to scan and can cause accidental taps.",
        "mobile-dashboard-navigation.png",
        "#6",
        "Recognition Rather Than Recall",
        "Minimize the user's memory load by making elements, actions, and options visible.",
        "Truncated labels force users to infer destinations from partial words and icons instead of recognizing them directly.",
        "Reduce mobile navigation to the highest-frequency destinations: Dashboard, Workouts, Nutrition, Progress, and More. Move Library, Live Workout, History, and Settings into a More screen or drawer. Also fix the underlying hydration/dev issue that produces the visible issue badge.",
    )

    add_mistake(
        doc,
        4,
        "Nutrition Logging Uses Technical Language and Unclear Loading Feedback",
        "Nutrition page — Food Diary, manual/OCR/food search tabs, diary content area",
        "The nutrition page provides useful logging modes, but the tab label \"OCR\" is technical and the main diary area shows \"Loading nutrition entries...\" while summary totals remain at zero. Users may not know whether food entries are still loading, missing, or unavailable.",
        "nutrition-logging.png",
        "#2",
        "Match Between System and the Real World",
        "The design should speak the users' language. Use words, phrases, and concepts familiar to the user.",
        "Most fitness users understand scanning a label, but may not understand OCR as a feature label.",
        "Rename OCR to Scan Label. After loading, show a meaningful diary empty state such as \"No foods logged for this day\" with an Add Food action. If loading fails, show a visible retry state instead of leaving totals and diary content ambiguous.",
    )

    add_mistake(
        doc,
        5,
        "Progress Charts and Filters Do Not Explain Empty Data",
        "Progress page — charts, time range controls, and personal records section",
        "The Progress & Insights screen shows large chart containers with little or no visible data, plus active time range controls and an Export button. Without explanation, users may interpret this as broken analytics rather than a normal empty-data state.",
        "progress-insights.png",
        "#5",
        "Error Prevention",
        "Good design carefully prevents problems from occurring in the first place.",
        "When charts are empty, the interface should prevent misinterpretation by explaining what data is required.",
        "Add empty chart states that explain the requirement: complete workouts with duration, sets, reps, and weight to generate charts and personal records. Disable or clarify filters and export until enough data exists, and add summary cards such as total workouts, total volume, average duration, and current streak.",
    )

    doc.save(OUT)
    print(OUT)
    return OUT


if __name__ == "__main__":
    build()
